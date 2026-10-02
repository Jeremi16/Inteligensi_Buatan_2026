"""Generate DOCX laporan Tugas Kelompok Kasus 01 (editable)."""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "TugasKelompok_Rx_nim1_nim2_nim3.docx"
NAVY = RGBColor(0x0F, 0x2A, 0x44)

doc = Document()

# base style
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(4)

sec = doc.sections[0]
sec.top_margin = Mm(15); sec.bottom_margin = Mm(15)
sec.left_margin = Mm(15); sec.right_margin = Mm(15)

def heading(text, size=16):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.color.rgb = NAVY; r.bold = True
    return p

def h2(text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.size = Pt(12); r.font.color.rgb = NAVY; r.bold = True
    return p

def para(text, bold=False, italic=False, size=11, justify=True):
    p = doc.add_paragraph()
    if justify: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    return p

def shade_cells(row, color="2F5AA0"):
    for c in row.cells:
        tcPr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), color)
        tcPr.append(shd)
        for par in c.paragraphs:
            for run in par.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.bold = True

def make_table(headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    shade_cells(t.rows[0])
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            t.rows[ri].cells[ci].text = str(val)
    for row in t.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in par.runs:
                    run.font.size = Pt(9)
    return t

# ---------- isi ----------
heading("Tugas Kelompok \u2013 Berbasis Kasus 01", 16)
para("Robot Kurir Kampus: Rute GU (Gerbang Utama) \u2192 LK (Lab Komputasi) \u2014 Metode: UCS, IDS, GBFS, A* \u2014 Inteligensi Buatan 2026")
para("Kelompok: Rx \u2014 Nama/NIM: [isi nim1, nim2, nim3] \u2014 File ini draf otomatis dari output program search_kampus.py. Ganti nama file menjadi TugasKelompok_Rx_nim1_nim2_nim3.docx/.pdf sesuai kelompok.",
     italic=True, size=9)

h2("1. Graf dan h(n)")
para("Graf undirected 12 simpul, 19 edge (terverifikasi dari gambar hires). Tie-break alfabetis agar deterministik.")
make_table(["Simpul", "Tetangga (bobot)"],
    [["GU","PB(3), R(4)"],["PB","GU(3), GKU(4), K(6)"],["R","GU(4), PR(3), M(5)"],
     ["GKU","PB(4), PR(2), K(3), IF(7)"],["K","PB(6), GKU(3), AS(5)"],
     ["PR","R(3), GKU(2), A(6), IF(4)"],["A","PR(6), M(3), LK(5)"],
     ["M","R(5), A(3), SC(6)"],["IF","GKU(7), PR(4), AS(4), LK(3)"],
     ["AS","K(5), IF(4), SC(6)"],["SC","M(6), AS(6), LK(4)"],["LK","A(5), IF(3), SC(4)"]])
doc.add_paragraph()
make_table(["Simpul","h(n)"],
    [["GU",12],["PB",10],["R",9],["GKU",7],["PR",6],["K",9],
     ["M",7],["A",4],["AS",5],["IF",2],["SC",3],["LK",0]])
para("Heuristik diuji admissible dan konsisten (h(n) \u2264 biaya aktual ke LK, dan h(n) \u2264 c(n,m)+h(m) untuk semua edge). Contoh: R(9)\u22643+6=9, PR(6)\u22644+2=6, GU(12)\u22644+9=13. Karena itu A* dengan h ini dijamin optimal pada graf ini.",
     size=9, italic=True)

h2("2. Hasil luaran program (verbatim)")
para("Perintah: python search_kampus.py. Salinan persis output yang dipakai sebagai bukti:", size=9, italic=True)
verbatim = """======================================================================
KASUS 01: Robot kurir kampus GU -> LK
======================================================================

[UCS] Urutan ekspansi (pop terkecil g):
  GU -> PB -> R -> GKU -> PR -> K -> M -> IF -> A -> AS -> LK
  g saat pop: GU=0, PB=3, R=4, GKU=7, PR=7, K=9, M=9, IF=11, A=12, AS=14, LK=14
  Rute : GU -> R -> PR -> IF -> LK
  Cost : 14 | node diekspansi: 11

[IDS] Iterative Deepening (DFS alfabetis, limit 0..n):
  limit=0 (1 kunjungan): GU => belum ketemu
  limit=1 (3 kunjungan): GU -> PB -> R => belum ketemu
  limit=2 (7 kunjungan): GU -> PB -> GKU -> K -> R -> M -> PR => belum ketemu
  limit=3 (17 kunjungan): GU -> PB -> GKU -> IF -> K -> PR -> K -> AS -> GKU -> R -> M -> A -> SC -> PR -> A -> GKU -> IF => belum ketemu
  limit=4 (6 kunjungan): GU -> PB -> GKU -> IF -> AS -> LK => KETEMU GU -> PB -> GKU -> IF -> LK
  Rute : GU -> PB -> GKU -> IF -> LK
  Cost (jumlah bobot) : 17 | depth solusi: 4 | total kunjungan kumulatif: 34

[GBFS] Urutan ekspansi (pop terkecil h):
  GU -> R -> PR -> IF -> LK
  Rute : GU -> R -> PR -> IF -> LK
  Cost : 14 | node diekspansi: 5

[A*] Urutan ekspansi (pop terkecil f=g+h):
  GU -> PB -> R -> PR -> IF -> GKU -> LK
  detail pop (g,h,f): GU(0,12,12), PB(3,10,13), R(4,9,13), PR(7,6,13), IF(11,2,13), GKU(7,7,14), LK(14,0,14)
  Rute : GU -> R -> PR -> IF -> LK
  Cost : 14 | node diekspansi: 7

--- RINGKASAN ---
UCS   | ekspansi: 11 | rute: GU -> R -> PR -> IF -> LK | cost: 14
IDS   | ekspansi: 34 (kumulatif, solusi depth 4) | rute: GU -> PB -> GKU -> IF -> LK | cost: 17
GBFS  | ekspansi: 5 | rute: GU -> R -> PR -> IF -> LK | cost: 14
A*    | ekspansi: 7 | rute: GU -> R -> PR -> IF -> LK | cost: 14"""
p = doc.add_paragraph()
p.style = doc.styles["Normal"]
r = p.add_run(verbatim)
r.font.name = "Consolas"; r.font.size = Pt(8)
para("Verifikasi optimalitas (brute force semua simple path): jalur termurah adalah GU-R-PR-IF-LK = 4+3+4+3 = 14. Runner-up GU-PB-GKU-PR-IF-LK = 3+4+2+4+3 = 16. Jadi cost 14 = optimal, cost 17 (IDS) = suboptimal.",
     size=9, italic=True)

h2("2b. Revisi: kenapa IDS tidak pilih jalur cost 16?")
para("Pertanyaan yang sering muncul: jalur GU-PB-GKU-PR-IF-LK = 3+4+2+4+3 = 16 jelas lebih murah dari rute IDS (17). Kenapa IDS tetap mengembalikan yang 17? Jawabannya ada di definisi IDS.")
make_table(["Jalur", "Perhitungan", "Cost", "Depth", "Keterangan"],
    [["GU-R-PR-IF-LK", "4+3+4+3", "14", "4", "Optimal (UCS/A*)"],
     ["GU-PB-GKU-PR-IF-LK", "3+4+2+4+3", "16", "5", "Lebih murah dari IDS, tapi depth 5"],
     ["GU-PB-GKU-IF-LK", "3+4+7+3", "17", "4", "Rute IDS: depth-4 pertama cabang PB"],
     ["GU-R-M-A-LK", "4+5+3+5", "17", "4", "Contoh depth-4 lain, cost sama 17"]])
para("IDS memakai limit depth (jumlah edge), bukan limit cost. Ia berhenti pada solusi depth terkecil yang ditemui DFS alfabetis pertama. Pada limit=4, cabang PB dieksplorasi sebelum cabang R, sehingga GU-PB-GKU-IF-LK (depth 4) ketemu duluan dan pencarian stop. Jadi 16 < 17 secara cost, tetapi 5 > 4 secara depth, maka IDS mengabaikannya. Jika ingin IDS yang sadar biaya, gunakan IDA* (limit f = g+h).",
     size=10)

h2("3. Tabel perbandingan")
make_table(["Aspek","UCS","IDS","GBFS","A*"],
    [["Prioritas","g(n)","depth (DFS)","h(n)","f=g+h"],
     ["Ekspansi*","11","34 kum. (6 di limit 4)","5","7"],
     ["Rute","GU-R-PR-IF-LK","GU-PB-GKU-IF-LK","GU-R-PR-IF-LK","GU-R-PR-IF-LK"],
     ["Cost","14 optimal","17 suboptimal","14 (kebetulan optimal)","14 optimal"],
     ["Edge","4","4","4","4"]])
para("*Ekspansi = jumlah node di-pop. IDS menghitung ulang tiap limit (total 1+3+7+17+6=34). Rute IDS dangkal (depth 4) tapi mahal karena meminimalkan jumlah edge, bukan bobot.", size=9, italic=True)

h2("4. Analisis perbandingan")
para("Urutan ekspansi. UCS melebar berdasar biaya: GU(0), PB(3), R(4), GKU(7), PR(7), K(9), M(9), IF(11), A(12), AS(14), LK(14) \u2014 11 node karena harus membuktikan tidak ada jalur <14. A* jauh lebih fokus: GU(12), PB(13), R(13), PR(13), IF(13), GKU(14), LK(14) \u2014 hanya 7 node, h memangkas cabang mahal (K, M, A, AS, SC tidak pernah di-pop). GBFS paling rakus: GU\u2192R\u2192PR\u2192IF\u2192LK (5 node), langsung mengikuti h menurun 12\u21929\u21926\u21922\u21920 tanpa menoleh ke cabang PB.")
para("Rute dan cost. UCS dan A* menemukan optimal 14 (GU-R-PR-IF-LK). GBFS di kasus ini juga 14, tapi itu kebetulan karena h mengarah ke koridor optimal; secara teori GBFS tidak menjamin optimal (ia mengabaikan g). IDS menemukan GU-PB-GKU-IF-LK = 3+4+7+3 = 17 pada depth 4. Depth sama-sama 4 dengan solusi optimal, tetapi IDS berhenti pada solusi depth-4 pertama yang ditemui DFS alfabetis (cabang PB lebih dulu dari R), bukan yang termurah. Jalur GU-PB-GKU-PR-IF-LK = 16 memang lebih murah dari 17, tetapi depth-nya 5 sehingga tidak dipertimbangkan IDS pada limit 4 (lihat Bab 2b). Ini menjawab instruksi 'gunakan bobot untuk IDS': bobot tidak mengubah urutan DFS, hanya dipakai menghitung cost akhir.")
para("Mengapa IDS boros? Setiap kenaikan limit mengulang seluruh pohon: limit 0,1,2,3 gagal (1+3+7+17 kunjungan) sebelum limit 4 berhasil. Total 34 kunjungan untuk graf 12 simpul. Kelebihannya hemat memori O(d) dan tetap lengkap; kekurangannya waktu berlipat dan buta biaya. Cocok bila edge seragam, tidak cocok di sini karena bobot bervariasi 2\u20137.")
para("Peran heuristik. Karena h admissible + konsisten (dicek per edge), A* optimal sekaligus efisien (7 vs 11 ekspansi). GBFS (5 ekspansi) paling cepat di sini tetapi rapuh: jika h overestimate di satu simpul, ia bisa tersesat (misal bila h(K) kecil, GBFS akan membelok ke K-AS-SC-LK yang mahal). UCS aman tanpa h tetapi membayar dengan ekspansi ekstra.")
para("Kesimpulan praktis untuk robot kurir. Untuk peta kampus ini: pakai A* (optimal + sedikit ekspansi). UCS sebagai baseline kebenaran. GBFS hanya bila butuh jawaban sangat cepat dan h terpercaya. IDS tidak disarankan untuk optimasi biaya; gunakan hanya bila memori sangat terbatas atau yang penting sedikit transit (edge), bukan jarak meter.")

h2("5. Cara reproduce & lampiran kode")
para("1) Simpan search_kampus.py dan DOCX ini dalam satu folder. 2) Jalankan python search_kampus.py \u2014 output harus sama persis dengan Bab 2. 3) Tie-break alfabetis, DLS memakai path-checking (hindari cycle satu jalur), UCS/GBFS/A* memakai graph-search. 4) Ganti placeholder Rx/NIM di cover, lalu export final menjadi TugasKelompok_Rx_nim1_nim2_nim3.pdf. Kode lengkap ada di file pendamping; ringkasan struktur: dict GRAPH + H, fungsi ucs(), dls()/ids(), gbfs(), astar(), path_cost().")
para("Draf ini belum final: tambahkan nama anggota, NIM, kelas, screenshot terminal, dan gambar graf kampus sebelum dikumpulkan.", size=9, italic=True)

doc.save(OUT)
print("OK ->", OUT)
